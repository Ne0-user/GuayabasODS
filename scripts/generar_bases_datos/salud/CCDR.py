from openpyxl import load_workbook
import sqlite3
import os

ruta_excel = os.path.join(os.path.dirname(__file__), "CCDR.xlsx")
wb = load_workbook(ruta_excel)
ws = wb.active

ruta_db = os.path.join(os.path.dirname(__file__), "..", "datos", "salud.db")
conn = sqlite3.connect(ruta_db)
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS CCDR (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    genero TEXT,
    estado TEXT,
    enfermedad TEXT,
    año INTEGER,
    muertes INTEGER,
    edad TEXT
)
""")

fila_inicio = 6
col_inicio = 2

col = col_inicio

while col <= ws.max_column:
    fila = fila_inicio

    if ws.cell(row=fila, column=col).value is None:
        break

    if ws.cell(row=2, column=col).value is not None:
        año = ws.cell(row=2, column=col).value

    if ws.cell(row=3, column=col).value is not None:
        enfermedad = ws.cell(row=3, column=col).value

    if ws.cell(row=4, column=col).value is not None:
        genero = ws.cell(row=4, column=col).value

    while True:
        valor = ws.cell(row=fila, column=col).value
        estado = ws.cell(row=fila, column=1).value
        edad = ws.cell(row=fila, column=2).value

        if valor is None:
            break

        if valor != "ND":
            print(f"Genero:{genero} Estado:{estado} Enfermedad:{enfermedad} Año:{año} Muertes:{valor} Edad:{edad}")

            cursor.execute("""
                INSERT INTO CCDR (genero, estado, enfermedad, año, muertes, edad)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (genero, estado, enfermedad, año, valor, edad))

        fila += 1
    col += 1

conn.commit()
conn.close()

print("Datos insertados correctamente ")
