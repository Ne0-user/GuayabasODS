from openpyxl import load_workbook
import sqlite3
import os

ruta_excel = os.path.join(os.path.dirname(__file__), "PEI.xlsx")
wb = load_workbook(ruta_excel)
ws = wb.active

ruta_db = os.path.join(os.path.dirname(__file__), "../../", "datos", "trabajo.db")
conn = sqlite3.connect(ruta_db)
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS PEI (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    genero TEXT,
    año INTEGER,
    porcentaje REAL
)
""")

fila_inicio = 4
col_inicio = 2

col = col_inicio

while col <= ws.max_column:
    fila = fila_inicio

    if ws.cell(row=fila, column=col).value is None:
        break

    if ws.cell(row=4, column=col).value is not None:
        genero = ws.cell(row=3, column=col).value

    while True:
        valor = ws.cell(row=fila, column=col).value
        ano = ws.cell(row=fila, column=1).value

        if valor is None:
            break

        if valor != "ND":
            print(f"Genero:{genero} Año:{ano} porcentaje:{valor}")

            cursor.execute("""
                INSERT INTO PEI (genero,año, porcentaje)
                VALUES (?, ?, ?)
            """, (genero, ano,valor))

        fila += 1
    col += 1

conn.commit()
conn.close()

print("Datos insertados correctamente ")
