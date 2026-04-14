# -------------------------------------------------------------------------------
# -------------------------------------------------------------------------------
# No estoy seguro de que funcionen al 100 ya que no los puedo ejecutar por pedos en el anaconda prompt
# IMPORTANTE: Corrige la ruta del archivo, tambien hay que cambiar los .xls de github por .xlsx
# Dales una revisada antes de ejecutarlos porfa
# -------------------------------------------------------------------------------
# -------------------------------------------------------------------------------

from openpyxl import load_workbook
import sqlite3
import os

ruta_excel = os.path.join(os.path.dirname(__file__), "8.5.2_sh_es.xlsx")
wb = load_workbook(ruta_excel)
ws = wb.active

ruta_db = os.path.join(os.path.dirname(__file__), "..", "datos", "salud.db")
conn = sqlite3.connect(ruta_db)
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS TDSE (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    año INTEGER,
    sexo TEXT,
    grupo_edad TEXT,
    estado TEXT,
    cantidad INTEGER
)
""")

fila_inicio = 6
col_inicio = 3

for col in range(col_inicio, ws.max_column + 1):

    año = ws.cell(row=2, column=col).value
    sexo = ws.cell(row=3, column=col).value
    grupo_edad = ws.cell(row=4, column=col).value

    for fila in range(fila_inicio, ws.max_row + 1):

        estado = ws.cell(row=fila, column=2).value
        valor = ws.cell(row=fila, column=col).value

        if estado is None:
            break

        if valor is None or valor == "ND":
            continue

        print(f"{año}, {sexo}, {grupo_edad}, {estado}, {valor}")

        cursor.execute("""
            INSERT INTO TDSE (año, sexo, grupo_edad, estado, cantidad)
            VALUES (?, ?, ?, ?, ?)
        """, (año, sexo, grupo_edad, estado, valor))

conn.commit()
conn.close()

print("Datos insertados correctamente")