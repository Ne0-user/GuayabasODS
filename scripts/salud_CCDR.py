from openpyxl import load_workbook
import sqlite3

wb=load_workbook("CCDR.xlsx")
ws=wb.active


fila_inicio = 6
col_inicio = 2

col = col_inicio

while col<=ws.max_column:
    fila = fila_inicio

    if ws.cell(row=fila, column=col).value is None:
       break

    if ws.cell(row=2,column=col).value is not None:
        año=ws.cell(row=2,column=col).value

    if ws.cell(row=3,column=col).value is not None:
        enfermedad=ws.cell(row=3,column=col).value
    
    if ws.cell(row=4,column=col).value is not None:
        genero=ws.cell(row=4,column=col).value

    while True:
        valor = ws.cell(row=fila, column=col).value
        estado= ws.cell(row=fila,column=1).value
        edad=ws.cell(row=fila,column=2).value

        if valor is None:
            break

        if valor != "ND":
            print(f"Genero:{genero} Estado:{estado} Enfermedad:{enfermedad} Año:{año} Muertes:{valor} Edad:{edad}")


        fila += 1
    col += 1

print(ws.cell(row=fila_inicio, column=col_inicio).value)
print(ws["B6"].value)
